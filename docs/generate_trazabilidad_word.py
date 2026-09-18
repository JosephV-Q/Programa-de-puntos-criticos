from docx import Document
from docx.shared import Pt

out = r"C:\Users\PC\Documents\Programa-de-puntos-criticos-main\docs\TRAZABILIDAD_IEEE830_WORD.docx"

doc = Document()

# Título principal
p = doc.add_paragraph()
r = p.add_run('Trazabilidad de requisitos IEEE830')
r.bold = True
r.font.size = Pt(22)

p = doc.add_paragraph()
r = p.add_run('Sistema para la Detección de Puntos Críticos de los Puentes')
r.font.size = Pt(15)

metadatos = [
    ('Estudiantes', 'Joseph David Vasquez Quintero, Daniel Eduardo Padilla Cerpa'),
    ('Institución', 'Institución Universitaria Tecnológica de Santander'),
    ('Asignatura', 'Nuevas Tecnologías'),
    ('Documento', 'Análisis de cumplimiento de requisitos frente al código fuente'),
    ('Fecha de análisis', '14/09/2026'),
]
for label, value in metadatos:
    p = doc.add_paragraph()
    p.add_run(f'{label}: ').bold = True
    p.add_run(value)

intro = (
    'Este documento relaciona los requisitos funcionales, no funcionales, historias de usuario y prioridades MoSCoW definidos '
    'en el documento IEEE830 con las partes concretas del proyecto. El análisis clasifica cada aspecto como Implementado, '
    'Parcial, No implementado o No verificado.'
)
doc.add_paragraph(intro)

# Sección de requisitos funcionales
p = doc.add_paragraph()
r = p.add_run('3. Requisitos funcionales')
r.bold = True
r.font.size = Pt(18)

func_rows = [
    ['ID', 'Nombre', 'Estado', 'Descripción', 'Evidencia principal'],
    ['RF01', 'Gestión de usuarios y acceso', 'Parcialmente implementado', 'Registro, login, JWT y control por roles.', 'backend/app/auth/router.py; frontend/src/components/LoginScreen.tsx'],
    ['RF02', 'Registro de información estructural', 'Parcialmente implementado', 'Captura de nombre, longitud, ancho, altura y material.', 'backend/app/schemas.py; frontend/src/AppShell.tsx'],
    ['RF03', 'Ejecución de simulaciones', 'Implementado como prototipo', 'Ejecución del cálculo con resultados y métricas.', 'backend/app/services/simulation_service.py; backend/app/routers.py'],
    ['RF04', 'Identificación y clasificación de zonas críticas', 'Implementado', 'Se clasifican zonas por severidad.', 'backend/app/services/simulation_service.py; frontend/src/components/BridgeScene3D.tsx'],
    ['RF05', 'Recomendación de puntos de monitoreo', 'Implementado en versión básica', 'Se recomiendan ubicaciones para sensores.', 'backend/app/services/simulation_service.py; frontend/src/components/BridgeScene3D.tsx'],
    ['RF06', 'Generación y exportación de información', 'Parcialmente implementado', 'Exportación JSON, CSV y HTML.', 'backend/app/routers.py; frontend/src/services/simulationService.ts'],
]
func_table = doc.add_table(rows=len(func_rows), cols=5)
func_table.style = 'Table Grid'
for i, row in enumerate(func_rows):
    cells = func_table.rows[i].cells
    for j, value in enumerate(row):
        cells[j].text = value

# Bloques detallados de RF
for code, title, state, desc, evidence in [
    ('RF01', 'Gestión de usuarios y acceso', 'Parcialmente implementado', 'El sistema permite registrar usuarios y autenticarlos con JWT. Se controla acceso por roles como admin, structural_engineer e instrumentation_specialist.', 'backend/app/auth/router.py; backend/app/auth/dependencies.py; frontend/src/components/UserPanel.tsx'),
    ('RF02', 'Registro de información estructural', 'Parcialmente implementado', 'Se almacena el modelo del puente con datos básicos, aunque falta una parte del alcance estructural completo.', 'backend/app/persistence_models.py; backend/app/schemas.py; frontend/src/AppShell.tsx'),
    ('RF03', 'Ejecución de simulaciones', 'Implementado como prototipo demostrativo', 'El sistema ejecuta la simulación y calcula desplazamiento, aceleración, zonas críticas y puntos de monitoreo.', 'backend/app/services/simulation_service.py; backend/app/routers.py; frontend/src/viewmodels/useSimulationViewModel.ts'),
    ('RF04', 'Identificación y clasificación de zonas críticas', 'Implementado', 'Se calcula la severidad de cada zona clasificándola como low, medium, high o critical.', 'backend/app/services/simulation_service.py; frontend/src/components/BridgeScene3D.tsx'),
    ('RF05', 'Recomendación de puntos de monitoreo', 'Implementado en versión básica', 'Se recomienda puntos de monitoreo en zonas de alta severidad, aunque aún no existe optimización avanzada.', 'backend/app/services/simulation_service.py; frontend/src/components/BridgeScene3D.tsx'),
    ('RF06', 'Generación y exportación de información', 'Parcialmente implementado', 'Se exportan resultados en JSON, CSV y HTML. Faltan formatos como PDF y reportes técnicos completos.', 'backend/app/routers.py; frontend/src/services/simulationService.ts; frontend/src/AppShell.tsx'),
]:
    doc.add_paragraph()
    p = doc.add_paragraph()
    r = p.add_run(f'{code}. {title}')
    r.bold = True
    r.font.size = Pt(12)
    doc.add_paragraph(f'Estado: {state}')
    doc.add_paragraph(f'Descripción: {desc}')
    doc.add_paragraph(f'Evidencia: {evidence}')

# Requisitos no funcionales
p = doc.add_page_break()
p = doc.add_paragraph()
r = p.add_run('4. Requisitos no funcionales')
r.bold = True
r.font.size = Pt(18)

nfr_rows = [
    ['ID', 'Nombre', 'Estado', 'Descripción', 'Observaciones'],
    ['RNF-R1', 'Tiempo de respuesta', 'No verificado', 'Se espera que la simulación responda rápidamente.', 'No hay mediciones ni pruebas de desempeño.'],
    ['RNF-R2', 'Estado durante simulaciones largas', 'Parcialmente implementado', 'El frontend muestra carga y bloqueo de interacción.', 'Faltan porcentaje de avance y cancelación.'],
    ['RNF-S1', 'Autenticación y control de acceso', 'Implementado', 'Se usa JWT y autorización por roles.', 'Protege endpoints y datos sensibles.'],
    ['RNF-S2', 'Protección de información', 'Parcialmente implementado', 'Las contraseñas no se almacenan en texto plano.', 'Falta secreto fuerte y auditoría en producción.'],
    ['RNF-F1', 'Integridad de resultados', 'Implementado', 'La simulación se guarda completa con parámetros, métricas y zonas.', 'Se confirma transacción y persistencia.'],
    ['RNF-D1', 'Disponibilidad del sistema', 'No implementado formalmente', 'Existe salud básica del sistema.', 'Falta monitoreo y alertas.'],
]

nfr_table = doc.add_table(rows=len(nfr_rows), cols=5)
nfr_table.style = 'Table Grid'
for i, row in enumerate(nfr_rows):
    cells = nfr_table.rows[i].cells
    for j, value in enumerate(row):
        cells[j].text = value

# Conclusión
p = doc.add_paragraph()
r = p.add_run('Conclusión')
r.bold = True
r.font.size = Pt(16)
doc.add_paragraph('El proyecto presenta una base funcional sólida para autenticar usuarios, ejecutar simulaciones y visualizar resultados. No obstante, la solución sigue siendo prototípica y requiere reforzar los aspectos de seguridad, trazabilidad, exportación avanzada y validación de requisitos no funcionales.')

# Ajustar tamaño de fuente en tablas
for table in [func_table, nfr_table]:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9)

# Guardar archivo final
try:
    doc.save(out)
    print(f'Archivo generado correctamente: {out}')
except Exception as exc:
    print(f'ERROR: {exc}')
    raise
